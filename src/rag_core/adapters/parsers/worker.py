"""Private subprocess entry point. Input paths are created by ParserRegistry only."""

import json
import re
import sys
from pathlib import Path
from typing import cast

from rag_core.contracts.v1 import DocxLocator, OffsetRange, PdfLocator
from rag_core.domain.documents import (
    Block,
    DocumentFormat,
    ParsedDocument,
    ParseError,
    ParserLimits,
    SourceIdentity,
)

from .archive import archive_preflight
from .office import csv_blocks, pptx_blocks, xlsx_blocks
from .tables import table_rows, table_text
from .text import SafeHTML, decode, text_blocks


def docx_blocks(path: Path) -> list[Block]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(str(path))
    blocks: list[Block] = []
    headings: list[tuple[int, str]] = []
    paragraph_index = 0
    table_index = 0

    def paragraph(item: Paragraph, part: str, role: str) -> None:
        nonlocal paragraph_index, headings
        index = paragraph_index
        paragraph_index += 1  # blank paragraphs still occupy their source position
        value = item.text
        if not value.strip():
            return
        kind = role
        style = item.style.name if item.style else ""
        match = re.fullmatch(r"Heading ([1-9])", style)
        if match and role == "paragraph":
            level = int(match[1])
            headings = [(n, title) for n, title in headings if n < level]
            headings.append((level, value))
            kind = "heading"
        blocks.append(
            Block(
                kind=kind,
                text=value,
                heading_path=tuple(title for _, title in headings),
                locator=DocxLocator(
                    kind="docx",
                    heading_path=[title for _, title in headings],
                    paragraph=index,
                    offsets=OffsetRange(start=0, end=len(value)),
                ),
                source_part=part,
                source_path=item._p.getroottree().getpath(item._p),
            )
        )

    def table(item: Table, part: str) -> None:
        nonlocal table_index
        index = table_index
        table_index += 1
        rows = table_rows(tuple(cell.text for cell in row.cells) for row in item.rows)
        value = table_text(rows)
        if value.strip():
            blocks.append(
                Block(
                    kind="table",
                    text=value,
                    rows=rows,
                    table_headers=rows[:1],
                    heading_path=tuple(title for _, title in headings),
                    locator=DocxLocator(
                        kind="docx",
                        heading_path=[title for _, title in headings],
                        table=index,
                        offsets=OffsetRange(start=0, end=len(value)),
                    ),
                    source_part=part,
                    source_path=item._tbl.getroottree().getpath(item._tbl),
                )
            )
        seen: set[str] = set()
        for row in item.rows:
            for cell in row.cells:
                cell_path = cell._tc.getroottree().getpath(cell._tc)
                if cell_path in seen:
                    continue  # merged cells can alias the same XML node
                seen.add(cell_path)
                for nested in cell.tables:
                    table(nested, part)

    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            paragraph(item, "word/document.xml", "paragraph")
        else:
            table(item, "word/document.xml")
    seen_parts: set[str] = set()
    for section in document.sections:
        for container, role in (
            (section.header, "header"),
            (section.first_page_header, "header"),
            (section.even_page_header, "header"),
            (section.footer, "footer"),
            (section.first_page_footer, "footer"),
            (section.even_page_footer, "footer"),
        ):
            if container.is_linked_to_previous:
                continue
            part = str(container.part.partname).lstrip("/")
            if part in seen_parts:
                continue
            seen_parts.add(part)
            headings = []
            for item in container.iter_inner_content():
                if isinstance(item, Paragraph):
                    paragraph(item, part, role)
                else:
                    table(item, part)
    return blocks


def pdf_blocks(path: Path, limits: ParserLimits) -> tuple[list[Block], int, tuple[int, ...]]:
    from docling_parse.pdf_parser import ContentConfig, ContentLevel, DoclingPdfParser
    from pypdf import PdfReader

    with path.open("rb") as stream:
        reader = PdfReader(stream, strict=True)
        if reader.is_encrypted:
            raise ParseError("encrypted_document")
        page_count = len(reader.pages)
        if page_count > limits.max_pages:
            raise ParseError("page_limit")
        labels = reader.page_labels
    parser = DoclingPdfParser(loglevel="fatal")
    document = parser.load(
        path,
        content_config=ContentConfig(
            char_cells_content_level=ContentLevel.SKIP,
            word_cells_content_level=ContentLevel.SKIP,
            line_cells_content_level=ContentLevel.COMPUTE_AND_MATERIALIZE,
            shapes_content_level=ContentLevel.SKIP,
            bitmaps_content_level=ContentLevel.SKIP,
            include_bitmap_bytes=False,
        ),
    )
    blocks: list[Block] = []
    empty: list[int] = []
    try:
        if document.number_of_pages() != page_count:
            raise ParseError("corrupt_document")
        for page_no, page in document.iterate_pages():
            offset = 0
            found = False
            for index, cell in enumerate(page.textline_cells):
                value = cell.text
                if value.strip():
                    found = True
                    rect = cell.rect.to_bounding_box()
                    blocks.append(
                        Block(
                            kind="paragraph",
                            text=value,
                            locator=PdfLocator(
                                kind="pdf",
                                page=page_no,
                                block=index,
                                printed_page_label=labels[page_no - 1],
                                offsets=OffsetRange(start=offset, end=offset + len(value)),
                            ),
                            bbox=(rect.l, rect.t, rect.r, rect.b),
                        )
                    )
                offset += len(value) + 1  # canonical page text = native cells joined by LF
            if not found:
                empty.append(page_no)
            enforce(blocks, limits)
    finally:
        document.unload()
    return blocks, page_count, tuple(empty)


def enforce(blocks: list[Block], limits: ParserLimits) -> None:
    if (
        len(blocks) > limits.max_blocks
        or sum(len(block.text) for block in blocks) > limits.max_text_chars
    ):
        raise ParseError("extraction_limit")


def parse_local(
    path: Path, format_: DocumentFormat, source: SourceIdentity, limits: ParserLimits
) -> ParsedDocument:
    with path.open("rb") as stream:
        signature = stream.read(8)
    page_count = None
    missing: tuple[int, ...] = ()
    warnings: tuple[str, ...] = ()
    if format_ == "pdf":
        if not signature.startswith(b"%PDF-"):
            raise ParseError("mime_mismatch")
        blocks, page_count, missing = pdf_blocks(path, limits)
        revision = "docling-parse-7.22.1/text-v1"
    elif format_ in {"docx", "xlsx", "pptx"}:
        if signature == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
            raise ParseError("encrypted_document")
        if not signature.startswith(b"PK"):
            raise ParseError("mime_mismatch")
        if archive_preflight(path, limits, format_):
            warnings = ("external_relationships_ignored",)
        if format_ == "docx":
            blocks = docx_blocks(path)
            revision = "python-docx-1.2.0/table-v2"
        elif format_ == "xlsx":
            blocks, office_warnings = xlsx_blocks(path, limits)
            warnings += office_warnings
            revision = "openpyxl-3.1.5/table-v1"
        else:
            blocks, office_warnings = pptx_blocks(path, limits)
            warnings += office_warnings
            revision = "python-pptx-1.0.2/table-v1"
    else:
        if signature.startswith((b"%PDF-", b"PK\x03\x04", b"\xd0\xcf\x11\xe0")):
            raise ParseError("mime_mismatch")
        text = decode(path.read_bytes())
        html = bool(
            re.search(
                r"<(?:!doctype\s+html|html|head|body|h[1-6]|p|div|table|section|article|pre|script)(?:\s|>)",
                text,
                re.I,
            )
        )
        if format_ == "csv":
            # HTML/formula-looking field contents are inert CSV source data.
            blocks = csv_blocks(text, limits)
            revision = "stdlib-csv/table-v1"
        elif format_ == "html":
            if not html:
                raise ParseError("mime_mismatch")
            blocks = SafeHTML(text).finish()
            revision = "stdlib-html/table-v2"
        else:
            # Markdown may contain raw HTML or HTML examples in fenced code.
            # Plain text/Markdown have no distinct magic: preserve such text as data.
            if format_ == "txt" and re.search(r"<(?:!doctype\s+html|html)(?:\s|>)", text, re.I):
                raise ParseError("mime_mismatch")
            blocks = text_blocks(text, format_ == "md")
            revision = "utf8-lines/table-v2" if format_ == "md" else "utf8-lines/text-v1"
    enforce(blocks, limits)
    if not blocks:
        raise ParseError("ocr_required" if format_ == "pdf" else "empty_extraction")
    return ParsedDocument(
        source=source,
        format=format_,
        parser_revision=revision,
        blocks=tuple(blocks),
        page_count=page_count,
        needs_ocr_pages=missing,
        quality="partial" if missing else "text",
        warnings=warnings,
    )


def main() -> int:
    request = Path(sys.argv[1])
    output = request.parent / "result.json"
    limits = ParserLimits()
    try:
        data = json.loads(request.read_text(encoding="utf-8"))
        limits = ParserLimits.model_validate(data["limits"])
        result = parse_local(
            request.parent / data["input"],
            cast(DocumentFormat, data["format"]),
            SourceIdentity.model_validate(data["source"]),
            limits,
        )
        payload = result.model_dump_json()
        if len(payload.encode("utf-8")) > limits.max_result_bytes:
            raise ParseError("extraction_limit")
    except ParseError as exc:
        payload = json.dumps({"error": exc.code})
    except Exception:
        payload = json.dumps({"error": "corrupt_document"})
    output.write_text(payload, encoding="utf-8")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
