"""Current-session SQL authority and original-source citation round trips, no live LLM."""

import hashlib
import json
import re
from dataclasses import replace
from pathlib import Path
from uuid import uuid4

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches
from pydantic import TypeAdapter
from pypdf import PdfReader
from sqlalchemy import text

from rag_core.adapters.parsers import ParserRegistry
from rag_core.adapters.parsers.registry import FORMATS
from rag_core.adapters.persistence.citations import PostgresCitationRepository
from rag_core.adapters.tokenizer import BgeM3Tokenizer
from rag_core.application.answers import AnswerAssembler
from rag_core.application.citations import CitationResolver
from rag_core.contracts.v1 import CitationResolveResponse, ErrorEnvelope, QueryResponse
from rag_core.domain.answers import AnswerError, source_citations
from rag_core.domain.chunking import Chunk, StructuralChunker
from rag_core.domain.documents import SourceIdentity
from rag_core.domain.evidence import EvidenceError
from rag_core.domain.ingestion import vector_unit_id
from rag_core.domain.metadata import ScopeError
from tests.fixtures.answer_support import good_answer, protocol_provider
from tests.fixtures.evidence_support import evidence_corpus
from tests.integration import conftest as pg_fixtures
from tests.integration.test_qdrant_scope import fixture
from tests.integration.test_retrieval import real_models
from tests.integration.test_text_parsers import pdf

pg_url = pg_fixtures.pg_url
__all__ = ["evidence_corpus", "fixture", "real_models"]
pytestmark = [pytest.mark.security, pytest.mark.asyncio]


def repositories(e):
    reader = PostgresCitationRepository(e.f.engine, e.f.sessions)
    return reader, CitationResolver(reader, e.f.sessions)


@pytest.mark.parametrize("foreign_index", [0, 1, 2])
async def test_resolver_retained_other_session_user_app_is_not_authorized(
    evidence_corpus, foreign_index
):
    e = evidence_corpus
    _, chunk, _ = e.foreign[foreign_index]
    reader, resolver = repositories(e)
    scope = await e.f.snapshot()
    with pytest.raises(ScopeError, match=r"^not_found$"):
        await reader.load(scope, (chunk.id,))
    with pytest.raises(ScopeError, match=r"^not_found$"):
        await resolver.resolve(e.f.owner, scope.session_id, chunk.id, uuid4())
    with pytest.raises(ScopeError, match=r"^not_found$"):
        await resolver.resolve(
            e.f.owner,
            scope.session_id,
            e.chunks["revenue"].id,
            uuid4(),
            document_ids=(e.chunks["capital-en"].source.document_id,),
        )


@pytest.mark.parametrize(
    "stage", ["before-prompt", "generation", "repair", "delete", "generation-change"]
)
async def test_scope_change_aborts_without_success_or_more_provider_calls(evidence_corpus, stage):
    e = evidence_corpus
    ctx, retrieval, selector = await e.query("Vietnam capital?", labels=("capital-en",))
    selected = await selector.select(ctx, retrieval)

    async def mutate(data):
        if stage == "delete":
            await e.f.sessions.delete_session(e.f.owner, e.f.session.session_id)
        elif stage == "generation-change":
            await e.f.publish(await e.f.generation(e.targets["capital-en"]))
        else:
            await e.f.sessions.detach_document(
                e.f.owner, ctx.scope.session_id, e.chunks["capital-en"].source.document_id
            )

    if stage == "before-prompt":
        await mutate({})

    def bad(data):
        return {"answer": "invented [c999]", "citations": []}

    async with protocol_provider(
        responses=(bad,) if stage == "repair" else (),
        hook=None if stage == "before-prompt" else mutate,
    ) as (llm, calls, _):
        reader, _ = repositories(e)
        with pytest.raises(ScopeError, match="session_scope_changed"):
            await AnswerAssembler(selector, reader, llm, e.tokenizer).assemble(
                ctx, selected, uuid4()
            )
        assert len(calls) == (0 if stage == "before-prompt" else 1)
        print(f"T24 actual PG {stage} invalidation -> no final/repair success")


async def test_durable_text_and_locator_tamper_during_provider_rejected(evidence_corpus):
    e = evidence_corpus
    ctx, retrieval, selector = await e.query("Vietnam capital?", labels=("capital-en",))
    selected = await selector.select(ctx, retrieval)

    async def corrupt(data):
        async with e.f.engine.begin() as conn:
            await conn.execute(
                text("UPDATE chunks SET text='FORGED OUTSIDE SOURCE' WHERE chunk_id=:id"),
                {"id": e.chunks["capital-en"].id},
            )

    async with protocol_provider(hook=corrupt) as (llm, calls, _):
        reader, _ = repositories(e)
        with pytest.raises(EvidenceError, match="evidence_mapping_mismatch"):
            await AnswerAssembler(selector, reader, llm, e.tokenizer).assemble(
                ctx, selected, uuid4()
            )
        assert len(calls) == 1


async def test_forged_selection_and_extra_metadata_are_rejected(evidence_corpus):
    e = evidence_corpus
    ctx, retrieval, selector = await e.query("Vietnam capital?", labels=("capital-en",))
    selected = await selector.select(ctx, retrieval)
    reader, _ = repositories(e)
    async with protocol_provider() as (llm, calls, _):
        fake = replace(selected, passages=(replace(selected.passages[0], chunk=e.foreign[0][1]),))
        with pytest.raises((ScopeError, EvidenceError, AnswerError)):
            await AnswerAssembler(selector, reader, llm, e.tokenizer).assemble(ctx, fake, uuid4())
        assert calls == []

    def injected(data):
        result = good_answer(data)
        result["citations"][0]["locator"] = {"kind": "pdf", "page": 999}
        return result

    async with protocol_provider(responses=(injected,)) as (llm, calls, _):
        with pytest.raises(AnswerError, match="invalid_citation"):
            await AnswerAssembler(selector, reader, llm, e.tokenizer).assemble(
                ctx, selected, uuid4()
            )
        assert len(calls) == 2


async def persist_original(f, path):
    target = await f.version()
    source = SourceIdentity(
        document_id=target.pair.document_id,
        version_id=target.pair.version_id,
        sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
    )
    parsed = ParserRegistry(path.parent / "sandbox").parse(
        path, filename=path.name, content_type=FORMATS[path.suffix][1], source=source
    )
    chunks = StructuralChunker(BgeM3Tokenizer()).chunk(
        parsed, generation_id=target.pair.generation_id
    )
    adapter = TypeAdapter(Chunk)
    async with f.engine.begin() as conn:
        await conn.execute(
            text("UPDATE documents SET filename=:name WHERE document_id=:doc"),
            {"name": path.name, "doc": source.document_id},
        )
        await conn.execute(
            text("UPDATE document_versions SET source_sha256=:sha WHERE version_id=:v"),
            {"sha": source.sha256, "v": source.version_id},
        )
        for c in chunks:
            await conn.execute(
                text("""
                INSERT INTO chunks (chunk_id,app_id,owner_id,version_id,generation_id,ordinal,
                    text,checksum,token_count,language,unit_id,source_map)
                VALUES (:id,:app,:owner,:version,:generation,:ordinal,:text,:sha,:tokens,'en',:unit,CAST(:map AS jsonb))
            """),
                dict(
                    id=c.id,
                    app=f.owner.app_id,
                    owner=f.owner.user_id,
                    version=source.version_id,
                    generation=c.generation_id,
                    ordinal=c.ordinal,
                    text=c.text,
                    sha=c.checksum,
                    tokens=c.token_count,
                    unit=vector_unit_id(c),
                    map=json.dumps(adapter.dump_python(c, mode="json", exclude={"text"})),
                ),
            )
    await f.publish(target)
    await f.attach(target)
    return target, chunks, parsed


@pytest.mark.parametrize("format_", ["pdf", "docx", "xlsx", "pptx", "md", "html"])
async def test_original_source_round_trip_and_stale_detached_chunk(fixture, tmp_path, format_):
    f = fixture
    path = tmp_path / ("report." + format_)
    if format_ == "pdf":
        pdf(path, [["First physical page."], ["Second physical page."], ["Third physical page."]])
    elif format_ == "docx":
        d = Document()
        d.add_heading("Revenue", 1)
        d.add_paragraph("Revenue is 120 million VND.")
        d.save(path)
    elif format_ == "xlsx":
        b = Workbook()
        b.active.title = "Revenue"
        b.active.append(["Company", "Amount (million VND)"])
        b.active.append(["Acme", 120])
        b.save(path)
    elif format_ == "pptx":
        deck = Presentation()
        for i in (1, 2):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_textbox(
                Inches(1), Inches(1), Inches(5), Inches(1)
            ).text = f"Revenue slide {i}."
        deck.save(path)
    elif format_ == "md":
        path.write_text("# **Revenue**\n\nNet **revenue** was 120 million VND.\n", encoding="utf-8")
    else:
        path.write_text(
            "<html><body><h1>Revenue</h1><p>Net <b>revenue</b> was 120 million VND.</p></body></html>",
            encoding="utf-8",
        )
    before = path.read_bytes()
    target, chunks, parsed = await persist_original(f, path)
    reader = PostgresCitationRepository(f.engine, f.sessions)
    resolver = CitationResolver(reader, f.sessions)
    sources = await reader.load(await f.snapshot(), tuple(c.id for c in chunks))
    locators = []
    for source in sources:
        canonical = source_citations(source)
        response = await resolver.resolve(f.owner, f.session.session_id, source.chunk.id, uuid4())
        assert CitationResolveResponse.model_validate_json(response.model_dump_json()) == response
        assert response.citation == canonical[0] and response.citation.filename == path.name
        for citation in canonical:
            locators.append(citation.locator)
            if format_ == "pdf":
                loc = citation.locator
                book = PdfReader(path)
                assert citation.quote in book.pages[loc.page - 1].extract_text()
                assert loc.printed_page_label == book.page_labels[loc.page - 1]
            elif format_ == "docx":
                assert citation.quote == Document(path).paragraphs[citation.locator.paragraph].text
                assert not hasattr(citation.locator, "page")
            elif format_ == "xlsx":
                loc = citation.locator
                assert citation.quote == str(load_workbook(path)[loc.sheet][loc.cell_range].value)
                assert ":" not in loc.cell_range  # one actual mapped cell
            elif format_ == "pptx":
                loc = citation.locator
                assert (
                    citation.quote
                    == Presentation(path).slides[loc.slide - 1].shapes[loc.shape].text
                )
            else:
                # Source offsets refer to original decoded characters, including
                # Windows CRLF; read_text's universal-newline translation changes them.
                raw = path.read_bytes().decode("utf-8")
                normalized = (
                    raw.replace("**", "") if format_ == "md" else re.sub(r"<[^>]*>", "", raw)
                )
                assert citation.quote.replace("**", "") in normalized
                loc = citation.locator
                assert 0 <= loc.offsets.start < loc.offsets.end <= len(raw)
                assert loc.offsets in [segment.locator.offsets for segment in source.chunk.segments]
    if format_ == "pdf":
        assert {loc.page for loc in locators} == {1, 2, 3}
    assert path.read_bytes() == before and all(c.source == parsed.source for c in chunks)
    await f.publish(await f.generation(target))
    with pytest.raises(ScopeError, match="not_found"):
        await resolver.resolve(f.owner, f.session.session_id, chunks[0].id, uuid4())
    await f.publish(target)
    await f.sessions.detach_document(f.owner, f.session.session_id, target.pair.document_id)
    with pytest.raises(ScopeError, match="no_session_documents"):
        await resolver.resolve(f.owner, f.session.session_id, chunks[0].id, uuid4())
    async with f.engine.connect() as conn:
        assert (
            await conn.execute(
                text("SELECT count(*) FROM chunks WHERE version_id=:v"),
                {"v": target.pair.version_id},
            )
        ).scalar_one() == len(chunks)
    print(
        f"T24 original {format_} -> parser/tokenizer/PG -> citation -> original PASS; stale/detached denied; chunks/source retained"
    )


async def test_runbook_json_and_error_examples_validate_with_schema():
    runbook = (Path(__file__).resolve().parents[2] / "RUNBOOK.md").read_text(encoding="utf-8")
    examples = re.findall(r"<!-- T24-example: (\w+) -->\s*```json\s*(.*?)\s*```", runbook, re.S)
    assert len(examples) == 4
    models = {"QueryResponse": QueryResponse, "ErrorEnvelope": ErrorEnvelope}
    for name, payload in examples:
        validated = models[name].model_validate_json(payload)
        assert models[name].model_validate_json(validated.model_dump_json()) == validated
    print("T24 RUNBOOK: 2 QueryResponse +2 ErrorEnvelope JSON examples validated")
