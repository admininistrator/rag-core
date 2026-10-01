"""Real native parsers, original EN/VI fixtures; no corpus QA, services or mocks."""

import hashlib
import threading
import time
import zipfile
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from uuid import UUID

import pytest
from docx import Document
from pypdf import PdfReader, PdfWriter
from pypdf.generic import (
    ArrayObject,
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
    NumberObject,
    TextStringObject,
)

from rag_core.adapters.parsers import ParserRegistry
from rag_core.adapters.parsers.registry import FORMATS
from rag_core.domain.documents import ParsedDocument, ParseError, ParserLimits, SourceIdentity

pytestmark = pytest.mark.integration
VI = "Doanh thu quý một đạt 120 triệu đồng."
EN = "First quarter revenue was 120 million VND."


def pdf(path: Path, pages: list[list[str]]) -> None:
    """Generate a valid PDF with a Unicode ToUnicode CMap, table columns and page labels."""
    writer = PdfWriter()
    chars = sorted(set("".join(text for page in pages for text in page)))
    cmap = DecodedStreamObject()
    mapping = "\n".join(f"<{ord(char):04X}> <{ord(char):04X}>" for char in chars)
    cmap.set_data(
        (
            "/CIDInit /ProcSet findresource begin 12 dict begin begincmap\n"
            "/CIDSystemInfo << /Registry (Adobe) /Ordering (UCS) /Supplement 0 >> def\n"
            "/CMapName /FixtureUnicode def /CMapType 2 def\n"
            "1 begincodespacerange <0000> <FFFF> endcodespacerange\n"
            f"{len(chars)} beginbfchar\n{mapping}\nendbfchar\n"
            "endcmap CMapName currentdict /CMap defineresource pop end end"
        ).encode("ascii")
    )
    descendant = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/CIDFontType2"),
            NameObject("/BaseFont"): NameObject("/FixtureFont"),
            NameObject("/DW"): NumberObject(500),
            NameObject("/CIDSystemInfo"): DictionaryObject(
                {
                    NameObject("/Registry"): TextStringObject("Adobe"),
                    NameObject("/Ordering"): TextStringObject("Identity"),
                    NameObject("/Supplement"): NumberObject(0),
                }
            ),
            NameObject("/FontDescriptor"): DictionaryObject(
                {
                    NameObject("/Type"): NameObject("/FontDescriptor"),
                    NameObject("/FontName"): NameObject("/FixtureFont"),
                    NameObject("/Flags"): NumberObject(4),
                    NameObject("/ItalicAngle"): NumberObject(0),
                    NameObject("/Ascent"): NumberObject(800),
                    NameObject("/Descent"): NumberObject(-200),
                    NameObject("/CapHeight"): NumberObject(700),
                    NameObject("/StemV"): NumberObject(80),
                    NameObject("/FontBBox"): ArrayObject(
                        [NumberObject(n) for n in (0, -200, 1000, 800)]
                    ),
                }
            ),
        }
    )
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type0"),
            NameObject("/BaseFont"): NameObject("/FixtureFont"),
            NameObject("/Encoding"): NameObject("/Identity-H"),
            NameObject("/DescendantFonts"): ArrayObject([writer._add_object(descendant)]),
            NameObject("/ToUnicode"): writer._add_object(cmap),
        }
    )
    for texts in pages:
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
        )
        commands = []
        for index, text in enumerate(texts):
            commands.append(
                f"BT /F1 12 Tf 1 0 0 1 50 {740 - index * 25} Tm "
                f"<{text.encode('utf-16-be').hex()}> Tj ET"
            )
        stream = DecodedStreamObject()
        stream.set_data("\n".join(commands).encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    writer.set_page_label(0, len(pages) - 1, style="/r")
    writer.write(path)


def identity(path: Path) -> SourceIdentity:
    return SourceIdentity(
        document_id=UUID(int=1),
        version_id=UUID(int=2),
        sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
    )


def parse(
    path: Path,
    *,
    limits: ParserLimits | None = None,
    content_type: str | None = None,
    source: SourceIdentity | None = None,
) -> ParsedDocument:
    original = path.read_bytes()
    temp = path.parent / "sandbox"
    registry = ParserRegistry(temp, limits)
    try:
        return registry.parse(
            path,
            filename=path.name,
            content_type=content_type or FORMATS[path.suffix][1],
            source=source or identity(path),
        )
    finally:
        assert path.read_bytes() == original
        assert not temp.exists() or list(temp.iterdir()) == []


def test_pdf_physical_pages_unicode_and_table_source(tmp_path: Path) -> None:
    path = tmp_path / "report.pdf"
    pages = [["BÁO CÁO / REPORT", VI], [EN, "Khoản mục | Amount (triệu đồng)", "Doanh thu | 120"]]
    pdf(path, pages)
    result = parse(path)
    assert result.page_count == 2 and result.quality == "text"
    assert VI in "\n".join(block.text for block in result.blocks)
    reader = PdfReader(path)
    for block in result.blocks:
        locator = block.locator
        assert locator.kind == "pdf" and locator.block is not None and locator.offsets
        assert locator.printed_page_label == ("i" if locator.page == 1 else "ii")
        assert block.text in reader.pages[locator.page - 1].extract_text()
        assert block.bbox and all(isinstance(coordinate, float) for coordinate in block.bbox)
        canonical = "\n".join(pages[locator.page - 1])
        assert canonical[locator.offsets.start : locator.offsets.end] == block.text
    assert all(
        block.locator.page == 2
        for block in result.blocks
        if "120" in block.text and "Doanh thu |" in block.text
    )
    print(
        "PASS real Docling PDF: 2 physical pages, EN/VI, table headers/units, native offsets/bbox, printed i/ii"
    )


def test_docx_paragraph_table_heading_headers_and_no_fake_pages(tmp_path: Path) -> None:
    path = tmp_path / "report.docx"
    original = Document()
    original.add_heading("Báo cáo tài chính", level=1)
    original.add_paragraph(VI)
    original.add_page_break()
    original.add_heading("Revenue", level=2)
    original.add_paragraph(EN)
    table = original.add_table(rows=2, cols=2)
    for row, values in zip(
        table.rows, (("Khoản mục", "Amount (triệu đồng)"), ("Doanh thu", "120")), strict=True
    ):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    original.sections[0].header.paragraphs[0].text = "Confidential / Nội bộ"
    original.sections[0].footer.paragraphs[0].text = "Report footer"
    original.save(path)
    result = parse(path)
    assert result.page_count is None
    body = Document(str(path))
    for block in result.blocks:
        locator = block.locator
        assert locator.kind == "docx" and "page" not in locator.model_dump()
        assert block.source_part and block.source_path and locator.offsets
        if block.source_part == "word/document.xml":
            if locator.paragraph is not None:
                assert body.paragraphs[locator.paragraph].text == block.text
            else:
                assert locator.table == 0
                assert block.rows == (("Khoản mục", "Amount (triệu đồng)"), ("Doanh thu", "120"))
                assert block.text == "\n".join(
                    "\t".join(cell.text for cell in row.cells) for row in body.tables[0].rows
                )
                assert locator.heading_path == ["Báo cáo tài chính", "Revenue"]
        assert block.text[locator.offsets.start : locator.offsets.end] == block.text
    assert any(block.kind == "header" and "Nội bộ" in block.text for block in result.blocks)
    assert any(block.kind == "footer" for block in result.blocks)
    print(
        "PASS real python-docx: page break, EN/VI paragraphs, table/unit, heading path, header/footer XML round-trip; no page"
    )


@pytest.mark.parametrize("extension", ["txt", "md"])
def test_unicode_text_line_and_source_offset_round_trip(tmp_path: Path, extension: str) -> None:
    path = tmp_path / ("report." + extension)
    text = "\ufeff# Báo cáo\r\n\r\n" + VI + "\r\n" + EN + "\r\n\r\n## Appendix\r\nnext\r\n"
    path.write_bytes(text.encode("utf-8"))
    result = parse(path)
    for block in result.blocks:
        locator = block.locator
        assert locator.kind in {"txt", "md"} and locator.offsets
        assert text[locator.offsets.start : locator.offsets.end] == block.text
        assert locator.line_start == text[: locator.offsets.start].count("\n") + 1
        assert locator.line_end == text[: locator.offsets.end].count("\n") + 1
    assert VI in "\n".join(block.text for block in result.blocks)
    if extension == "md":
        assert result.blocks[-1].heading_path == ("Báo cáo", "Appendix")
    print(f"PASS real {extension}: exact UTF-8/BOM/CRLF Unicode offsets and one-based lines, EN/VI")


def test_html_heading_table_entities_and_raw_source_spans(tmp_path: Path) -> None:
    path = tmp_path / "report.html"
    text = (
        f"<html><body><h1>Báo cáo</h1><p>{VI} &amp; {EN}</p><table>"
        "<tr><th>Khoản mục</th><th>Amount (VND)</th></tr>"
        "<tr><td>Doanh thu</td><td>120</td></tr></table></body></html>"
    )
    path.write_text(text, encoding="utf-8")
    result = parse(path)
    paragraph = next(block for block in result.blocks if block.kind == "paragraph")
    assert paragraph.text == VI + " & " + EN
    for block in result.blocks:
        locator = block.locator
        assert locator.kind == "html" and locator.offsets and locator.heading_path == ["Báo cáo"]
        span = text[locator.offsets.start : locator.offsets.end]
        assert "Báo cáo" in span or VI in span or "Amount (VND)" in span
    assert result.blocks[-1].rows == (("Khoản mục", "Amount (VND)"), ("Doanh thu", "120"))
    print("PASS real HTML parser: EN/VI, entities, headings, table/unit and original HTML spans")


@pytest.mark.parametrize("extension", ["pdf", "docx"])
def test_corrupt_files_are_errors(tmp_path: Path, extension: str) -> None:
    path = tmp_path / ("broken." + extension)
    path.write_bytes(b"%PDF-1.7\nbroken\n%%EOF" if extension == "pdf" else b"PK\x03\x04bad-archive")
    with pytest.raises(ParseError, match=r"^corrupt_document$"):
        parse(path)


def test_real_encrypted_pdf_is_error(tmp_path: Path) -> None:
    path = tmp_path / "encrypted.pdf"
    writer = PdfWriter()
    writer.add_blank_page(612, 792)
    writer.encrypt("fixture-only-password", algorithm="AES-256")
    writer.write(path)
    with pytest.raises(ParseError, match=r"^encrypted_document$"):
        parse(path)
    print(
        "PASS real corrupt PDF/DOCX and AES-256 encrypted PDF: safe explicit errors; sources unchanged, temp empty"
    )


def test_safety_mime_size_hash_output_and_page_limits(tmp_path: Path) -> None:
    path = tmp_path / "report.txt"
    path.write_text(VI + "\n\n" + EN, encoding="utf-8")
    cases = [
        (ParserLimits(max_bytes=10), "source_too_large"),
        (ParserLimits(max_blocks=1), "extraction_limit"),
        (ParserLimits(max_text_chars=10), "extraction_limit"),
        (ParserLimits(max_result_bytes=10), "extraction_limit"),
    ]
    for limits, error in cases:
        with pytest.raises(ParseError, match=r"^" + error + "$"):
            parse(path, limits=limits)
    with pytest.raises(ParseError, match=r"^source_changed$"):
        parse(path, source=identity(path).model_copy(update={"sha256": "0" * 64}))
    with pytest.raises(ParseError, match=r"^mime_mismatch$"):
        parse(path, content_type="application/pdf")
    disguised = tmp_path / "disguised.txt"
    disguised.write_bytes(b"%PDF-1.7\n")
    with pytest.raises(ParseError, match=r"^mime_mismatch$"):
        parse(disguised)
    unsupported = tmp_path / "legacy.doc"
    unsupported.write_text("data", encoding="utf-8")
    with pytest.raises(ParseError, match=r"^unsupported_format$"):
        parse(unsupported, content_type="application/msword")
    report = tmp_path / "pages.pdf"
    pdf(report, [[EN], [VI]])
    with pytest.raises(ParseError, match=r"^page_limit$"):
        parse(report, limits=ParserLimits(max_pages=1))
    print(
        "PASS MIME signature/type/extension, byte/page/block/text/result limits, SHA mismatch, cleanup"
    )


def test_safety_timeout_kills_real_parser_and_cleans_temp(tmp_path: Path) -> None:
    path = tmp_path / "report.txt"
    path.write_text(VI, encoding="utf-8")
    started = time.monotonic()
    with pytest.raises(ParseError, match=r"^parser_timeout$"):
        parse(path, limits=ParserLimits(timeout_seconds=0.1))
    assert time.monotonic() - started < 5
    # A fresh parse proves the registry/temp source remains usable after cancellation.
    assert parse(path).blocks[0].text == VI
    print(
        "PASS real subprocess deadline: killed/reaped, sandbox removed, subsequent real parse succeeds"
    )


def test_safety_html_never_executes_or_fetches_external_refs(tmp_path: Path) -> None:
    hits: list[str] = []

    class Handler(BaseHTTPRequestHandler):
        def do_GET(self) -> None:
            hits.append(self.path)
            self.send_response(200)
            self.end_headers()

        def log_message(self, format: str, *args: object) -> None:
            pass

    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    url = f"http://127.0.0.1:{server.server_port}/external"
    path = tmp_path / "unsafe.html"
    path.write_text(
        f"<html><head><link rel='stylesheet' href='{url}'><style>HIDDEN_STYLE</style>"
        f"<script src='{url}'>fetch('{url}'); HIDDEN_SCRIPT</script></head><body>"
        f"<p onclick=\"fetch('{url}')\">{VI}</p><img src='{url}'>"
        f"<iframe src='{url}'>HIDDEN_FRAME</iframe><a href='{url}'>Visible link</a>"
        "<template><p>HIDDEN_TEMPLATE</p></template></body></html>",
        encoding="utf-8",
    )
    try:
        result = parse(path)
        text = "\n".join(block.text for block in result.blocks)
        assert VI in text and "Visible link" in text and "HIDDEN" not in text
        assert hits == []
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)
    print(
        "PASS actual HTTP canary: zero requests from script/style/image/iframe/link/event refs; hidden text excluded"
    )


def test_safety_docx_archive_limits_traversal_macros_entities(tmp_path: Path) -> None:
    path = tmp_path / "report.docx"
    original = Document()
    original.add_paragraph(VI)
    original.save(path)
    with pytest.raises(ParseError, match=r"^archive_limit$"):
        parse(path, limits=ParserLimits(max_archive_entries=1))
    with pytest.raises(ParseError, match=r"^archive_limit$"):
        parse(path, limits=ParserLimits(max_expanded_bytes=10))
    for name, body, expected in (
        ("../escape", b"bad", "unsafe_archive"),
        ("word/vbaProject.bin", b"macro", "unsafe_archive"),
        ("word/bomb.bin", b"x" * 100_000, "archive_limit"),
        (
            "word/dtd.xml",
            b'<!DOCTYPE x SYSTEM "http://127.0.0.1:1/no-fetch"><x/>',
            "corrupt_document",
        ),
        (
            "word/entity.xml",
            b'<!DOCTYPE x [<!ENTITY e SYSTEM "file:///private">]><x>&e;</x>',
            "corrupt_document",
        ),
    ):
        malicious = tmp_path / (expected + name.replace("/", "_").replace("..", "") + ".docx")
        malicious.write_bytes(path.read_bytes())
        with zipfile.ZipFile(malicious, "a", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr(name, body)
        with pytest.raises(ParseError, match=r"^" + expected + "$"):
            parse(malicious)
    assert not (tmp_path / "escape").exists()
    print(
        "PASS real ZIP preflight: entry/expanded/ratio limits, traversal, VBA, XML entity rejected without extraction"
    )


def test_safety_empty_and_mixed_pdf_require_explicit_ocr(tmp_path: Path) -> None:
    text = tmp_path / "empty.txt"
    text.write_text(" \n\t", encoding="utf-8")
    with pytest.raises(ParseError, match=r"^empty_extraction$"):
        parse(text)
    image = tmp_path / "blank.pdf"
    pdf(image, [[]])
    with pytest.raises(ParseError, match=r"^ocr_required$"):
        parse(image)
    mixed = tmp_path / "mixed.pdf"
    pdf(mixed, [[VI], []])
    result = parse(mixed)
    assert result.quality == "partial" and result.needs_ocr_pages == (2,)
    assert result.page_count == 2 and len(result.blocks) == 1
    print(
        "PASS empty extraction errors; native PDF missing text -> OCR required/partial page 2, no scan verification claimed"
    )


def test_safety_docx_external_relationships_are_data_only(tmp_path: Path) -> None:
    path = tmp_path / "external.docx"
    document = Document()
    document.add_paragraph(VI)
    document.save(path)
    with zipfile.ZipFile(path, "a") as archive:
        archive.writestr(
            "word/_rels/external.xml.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rExternal" Type="external" TargetMode="External" '
            'Target="http://127.0.0.1:1/no-fetch"/></Relationships>',
        )
    result = parse(path)
    assert result.warnings == ("external_relationships_ignored",)
    assert result.blocks[0].text == VI


def test_markdown_code_and_table_are_preserved_as_source_data(tmp_path: Path) -> None:
    path = tmp_path / "table.md"
    text = "# Báo cáo\n\n| Khoản mục | VND |\n| --- | --- |\n| Doanh thu | 120 |\n\n```python\n# not a heading\nprint('never executed')\n```\n"
    path.write_text(text, encoding="utf-8")
    result = parse(path)
    assert [block.kind for block in result.blocks] == ["heading", "table", "code"]
    assert result.blocks[1].rows == (("Khoản mục", "VND"), ("Doanh thu", "120"))
    assert result.blocks[2].heading_path == ("Báo cáo",)


def test_markdown_setext_and_fenced_delimiters_keep_heading_context(tmp_path: Path) -> None:
    path = tmp_path / "headings.md"
    path.write_text(
        "Báo cáo\n=======\n\n" + VI + "\n\n```\n---\n# code\n<p>data only</p>\n```\n",
        encoding="utf-8",
    )
    result = parse(path)
    assert [block.kind for block in result.blocks] == ["heading", "paragraph", "code"]
    assert all(block.heading_path == ("Báo cáo",) for block in result.blocks)


def test_safety_actual_format_encoding_and_docx_package_validation(tmp_path: Path) -> None:
    text = tmp_path / "encoding.txt"
    text.write_bytes(b"\xff\xfeinvalid-utf8")
    with pytest.raises(ParseError, match=r"^invalid_encoding$"):
        parse(text)
    text.write_bytes(b"text\x00binary")
    with pytest.raises(ParseError, match=r"^mime_mismatch$"):
        parse(text)
    disguised = tmp_path / "spreadsheet.docx"
    with zipfile.ZipFile(disguised, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("xl/workbook.xml", "<workbook/>")
    with pytest.raises(ParseError, match=r"^mime_mismatch$"):
        parse(disguised)
    html = tmp_path / "plain.html"
    html.write_text("ordinary text", encoding="utf-8")
    with pytest.raises(ParseError, match=r"^mime_mismatch$"):
        parse(html)


def test_docx_nested_table_and_blank_paragraph_source_paths(tmp_path: Path) -> None:
    from lxml import etree

    path = tmp_path / "nested.docx"
    document = Document()
    document.add_paragraph("")
    document.add_heading("Báo cáo", level=1)
    document.add_paragraph(VI)
    outer = document.add_table(rows=1, cols=1)
    outer.cell(0, 0).text = "Outer header"
    nested = outer.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested value (VND)"
    document.save(path)
    result = parse(path)
    with zipfile.ZipFile(path) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    tables = [block for block in result.blocks if block.kind == "table"]
    assert [block.locator.table for block in tables] == [0, 1]
    for block in result.blocks:
        assert block.source_path and len(root.xpath(block.source_path, namespaces=root.nsmap)) == 1
    assert result.blocks[1].text == VI and result.blocks[1].locator.paragraph == 2


def test_safety_registry_bounds_concurrent_real_execution(tmp_path: Path) -> None:
    from concurrent.futures import ThreadPoolExecutor

    path = tmp_path / "report.txt"
    path.write_text(VI, encoding="utf-8")
    registry = ParserRegistry(tmp_path / "sandbox")
    with ThreadPoolExecutor(max_workers=1) as pool:
        future = pool.submit(
            registry.parse,
            path,
            filename=path.name,
            content_type="text/plain",
            source=identity(path),
        )
        deadline = time.monotonic() + 5
        while not list(registry.temp_root.glob("parse-*/request.json")):
            assert time.monotonic() < deadline and not future.done()
            time.sleep(0.001)
        with pytest.raises(ParseError, match=r"^parser_busy$"):
            registry.parse(
                path, filename=path.name, content_type="text/plain", source=identity(path)
            )
        assert future.result(timeout=10).blocks[0].text == VI
    assert list(registry.temp_root.iterdir()) == []
