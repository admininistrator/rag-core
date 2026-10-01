"""Real parser -> real tokenizer -> chunks -> independently reopened original sources."""

import hashlib
import runpy
import zipfile
from pathlib import Path
from uuid import UUID

import pytest
from docx import Document
from lxml import etree
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfReader

from rag_core.adapters.parsers import ParserRegistry
from rag_core.adapters.parsers.registry import FORMATS
from rag_core.adapters.tokenizer import BgeM3Tokenizer
from rag_core.domain.chunking import ChunkingError, ChunkProfile, StructuralChunker
from rag_core.domain.documents import ParsedDocument, SourceIdentity

pytestmark = pytest.mark.integration
GENERATION = UUID(int=3)


@pytest.fixture(scope="module")
def chunker():
    return StructuralChunker(BgeM3Tokenizer(), ChunkProfile(max_tokens=96, overlap_tokens=12))


def parse(path: Path) -> ParsedDocument:
    data = path.read_bytes()
    source = SourceIdentity(
        document_id=UUID(int=1), version_id=UUID(int=2), sha256=hashlib.sha256(data).hexdigest()
    )
    sandbox = path.parent / "sandbox"
    parsed = ParserRegistry(sandbox).parse(
        path, filename=path.name, content_type=FORMATS[path.suffix][1], source=source
    )
    assert path.read_bytes() == data and list(sandbox.iterdir()) == []
    return parsed


def test_pdf_physical_pages_printed_labels_and_overlap_round_trip(tmp_path, chunker):
    path = tmp_path / "report.pdf"
    # Reuse actual Unicode PDF fixture generator from T13, no parser mocking.
    fixture = runpy.run_path(str(Path(__file__).with_name("test_text_parsers.py")))
    lines = [
        [f"Page {page} row {i}: Doanh thu 120 triệu đồng. " * 6 for i in range(4)]
        for page in (1, 2, 3)
    ]
    fixture["pdf"](path, lines)
    parsed = parse(path)
    original = PdfReader(path)
    chunks = chunker.chunk(parsed, generation_id=GENERATION)
    seen = set()
    for chunk in chunks:
        pages = {s.locator.page for s in chunk.segments}
        assert len(pages) == 1 and chunk.token_count <= 96
        for q in chunk.quotes(parsed, 0, len(chunk.text)):
            loc = q.segment.locator
            seen.add(loc.page)
            assert loc.printed_page_label == original.page_labels[loc.page - 1]
            assert q.text in original.pages[loc.page - 1].extract_text()
            block = parsed.blocks[q.segment.block_index]
            base = block.locator.offsets.start
            assert block.text[loc.offsets.start - base : loc.offsets.end - base] == q.text
    assert seen == {1, 2, 3}
    assert any(s.role == "overlap" for c in chunks for s in c.segments)
    print(f"PDF: physical pages={sorted(seen)}, chunks={len(chunks)}, printed=i/ii/iii")


def test_docx_headings_paragraphs_two_tables_and_xml_source(tmp_path, chunker):
    path = tmp_path / "report.docx"
    document = Document()
    document.add_heading("Doanh thu / Revenue", 1)
    document.add_paragraph("Doanh thu 120 triệu đồng. " * 100)
    for unit in ("VND", "USD"):
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Quarter"
        table.cell(0, 1).text = f"Revenue ({unit})"
        for i in range(30):
            cells = table.add_row().cells
            cells[0].text, cells[1].text = f"Q{i + 1}", f"{i * 120}"
    document.add_heading("Chi phí / Expenses", 1)
    document.add_paragraph("Cost 45 USD.")
    document.save(path)
    parsed = parse(path)
    chunks = chunker.chunk(parsed, generation_id=GENERATION)
    original = Document(path)
    seen_tables = set()
    for chunk in chunks:
        tables = {s.locator.table for s in chunk.segments if s.locator.table is not None}
        assert len(tables) <= 1
        for q in chunk.quotes(parsed, 0, len(chunk.text)):
            s, loc = q.segment, q.segment.locator
            assert not hasattr(loc, "page") and loc.heading_path
            if s.row is not None:
                seen_tables.add(loc.table)
                assert original.tables[loc.table].cell(s.row, s.column).text == q.text
            else:
                assert q.text in original.paragraphs[loc.paragraph].text
            with zipfile.ZipFile(path) as package:
                root = etree.fromstring(package.read(s.source_part))
            assert root.xpath(s.source_path, namespaces=root.nsmap)
        if tables:
            table_index = next(iter(tables))
            assert f"Revenue ({'VND' if table_index == 0 else 'USD'})" in chunk.text
    assert seen_tables == {0, 1}
    print(f"DOCX: tables={sorted(seen_tables)}, chunks={len(chunks)}, no invented pages")


def test_xlsx_regions_merged_context_whole_cells_formula_and_units(tmp_path, chunker):
    path = tmp_path / "report.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Revenue VND"
    sheet.append(["Báo cáo", None, None])
    sheet.merge_cells("A1:C1")
    sheet.append(["Quarter", "Amount (triệu đồng)", "Rate"])
    for i in range(30):
        sheet.append([f"Q{i + 1}", 120 * i, 0.25])
        sheet.cell(i + 3, 3).number_format = "0%"
    sheet.append([None, None, None])
    sheet.append(["Item", "Cost (USD)", "Rate"])
    sheet.append(["Hosting", 45, "=1+1"])
    other = book.create_sheet("Other USD")
    other.append(["Item", "Amount (USD)"])
    other.append(["Revenue", 900])
    book.save(path)
    parsed = parse(path)
    original = load_workbook(path, data_only=False)
    chunks = chunker.chunk(parsed, generation_id=GENERATION)
    seen_cells = set()
    grouped = False
    for chunk in chunks:
        assert len({s.locator.sheet for s in chunk.segments}) == 1
        grouped |= len({s.block_index for s in chunk.segments if s.role == "content"}) > 1
        for q in chunk.quotes(parsed, 0, len(chunk.text)):
            loc = q.segment.locator
            assert ":" not in loc.cell_range and not hasattr(loc, "page")
            cell = original[loc.sheet][loc.cell_range]
            seen_cells.add((loc.sheet, loc.cell_range))
            if cell.data_type == "f":
                assert q.text == "[formula cache unavailable]"
            elif isinstance(cell.value, (float, int)):
                assert float(q.text) == cell.value
            else:
                assert q.text == cell.value
        if "Hosting" in chunk.text:
            assert "Cost (USD)" in chunk.text and "triệu đồng" not in chunk.text
            assert "not recalculated" in chunk.text and "=1+1" not in chunk.text
        if any(c.number_format == "0%" for c in chunk.cells):
            assert "number format 0%" in chunk.text
        if any(
            s.role == "content"
            and s.locator.sheet == "Revenue VND"
            and s.locator.cell_range.startswith("B")
            and 3 <= int(s.locator.cell_range[1:]) <= 32
            for s in chunk.segments
        ):
            assert "Amount (triệu đồng)" in chunk.text and "Báo cáo" in chunk.text
    assert grouped and ("Revenue VND", "A1") in seen_cells
    assert ("Revenue VND", "C35") in seen_cells
    assert ("Other USD", "B2") in seen_cells
    print(
        f"XLSX: chunks={len(chunks)}, mapped cells={len(seen_cells)}, merged/header/units retained"
    )


def test_pptx_slide_shape_table_and_original_xml(tmp_path, chunker):
    path = tmp_path / "report.pptx"
    deck = Presentation()
    for i in range(2):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = (
            f"Slide {i + 1}: Doanh thu 120 triệu đồng. " * 80
        )
        table = slide.shapes.add_table(21, 2, Inches(1), Inches(3), Inches(5), Inches(2)).table
        table.cell(0, 0).text, table.cell(0, 1).text = "Quarter", "Amount (USD)"
        for row in range(1, 21):
            table.cell(row, 0).text, table.cell(row, 1).text = f"Q{row}", str(row * 120)
    deck.save(path)
    parsed = parse(path)
    chunks = chunker.chunk(parsed, generation_id=GENERATION)
    original = Presentation(path)
    seen = set()
    for chunk in chunks:
        assert len({(s.locator.slide, s.locator.shape) for s in chunk.segments}) == 1
        for q in chunk.quotes(parsed, 0, len(chunk.text)):
            s, loc = q.segment, q.segment.locator
            seen.add(loc.slide)
            shape = original.slides[loc.slide - 1].shapes[loc.shape]
            if s.row is not None:
                assert shape.table.cell(s.row, s.column).text == q.text
            else:
                assert q.text in shape.text
            with zipfile.ZipFile(path) as package:
                root = etree.fromstring(package.read(s.source_part))
            assert root.xpath(s.source_path, namespaces=root.nsmap)
    assert seen == {1, 2}
    print(f"PPTX: slides={sorted(seen)}, chunks={len(chunks)}, shape/table boundaries preserved")


@pytest.mark.parametrize(
    "extension,content",
    [
        ("txt", "Doanh thu 120 triệu đồng.\r\n" * 100),
        ("md", "# Report\n\n| Quarter | Amount (VND) |\n| --- | --- |\n| Q1 | 120 |\n"),
        (
            "html",
            "<h1>Report</h1><p>Doanh thu &amp; chi phí.</p><table><tr><th>Amount (VND)</th>"
            "</tr><tr><td>120</td></tr></table>",
        ),
        ("csv", 'Quarter,Amount (VND)\n"Q1\nEN/VI",120\nQ2,240\n'),
    ],
    ids=["txt", "md", "html", "csv"],
)
def test_other_text_and_table_formats_keep_normalized_quote_and_original_span(
    tmp_path, chunker, extension, content
):
    path = tmp_path / f"report.{extension}"
    path.write_bytes(content.encode("utf-8"))
    parsed = parse(path)
    for chunk in chunker.chunk(parsed, generation_id=GENERATION):
        for q in chunk.quotes(parsed, 0, len(chunk.text)):
            assert q.text and not hasattr(q.segment.locator, "page")
            block = parsed.blocks[q.segment.block_index]
            if q.segment.row is None:
                assert q.text in block.text
            else:
                assert q.text == block.rows[q.segment.row][q.segment.column]
            if extension == "html":
                loc = q.segment.locator
                assert content[loc.offsets.start : loc.offsets.end]


def test_original_table_unit_too_large_explicit_error(tmp_path, chunker):
    path = tmp_path / "long.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "Amount (VND)"
    table.cell(1, 0).text = "Doanh thu " * 300
    doc.save(path)
    with pytest.raises(ChunkingError, match="table_row_too_large"):
        chunker.chunk(parse(path), generation_id=GENERATION)


def test_xlsx_header_number_format_metadata_is_preserved_and_not_a_quote(tmp_path, chunker):
    path = tmp_path / "numeric-header.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append([0.25, "Amount (VND)"])
    sheet["A1"].number_format = "0%"
    sheet.append([1, 120])
    book.save(path)
    parsed = parse(path)
    chunks = chunker.chunk(parsed, generation_id=GENERATION)
    for chunk in chunks:
        assert "A1: number format 0%" in chunk.text
        assert any(c.reference == "A1" and c.number_format == "0%" for c in chunk.cells)
        start = chunk.text.index("A1: number format 0%")
        with pytest.raises(ChunkingError, match="quote_has_no_source"):
            chunk.quotes(parsed, start, start + len("A1: number format 0%"))
