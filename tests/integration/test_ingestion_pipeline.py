"""Real PG/MinIO/Redis/Qdrant/shared BGE-M3/CPU OCR; no fake dependency success."""

import asyncio
import hashlib
import json
import os
import signal
import subprocess
import threading
import time
from collections.abc import AsyncIterator, Callable
from contextlib import suppress
from dataclasses import dataclass
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any
from uuid import UUID, uuid4

import boto3
import httpx
import pytest
import pytest_asyncio
import redis
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from qdrant_client import AsyncQdrantClient
from sqlalchemy import text
from sqlalchemy.engine import URL
from sqlalchemy.ext.asyncio import AsyncEngine

from rag_core.adapters.broker.celery import CeleryJobPublisher
from rag_core.adapters.persistence.database import build_engine
from rag_core.adapters.persistence.ingestion import PostgresIngestionRepository
from rag_core.adapters.persistence.registrations import (
    OutboxDispatcher,
    PostgresRegistrationRepository,
)
from rag_core.adapters.storage.s3 import S3StorageReader, StorageLocation, StorageRegistry
from rag_core.auth import Principal
from rag_core.contracts.v1 import DocumentRegisterRequest
from rag_core.domain.ingestion import IngestionError
from rag_core.domain.metadata import ScopeError, VersionGeneration
from rag_core.domain.registration import RegisteredUpload
from rag_core.domain.vectors import GenerationScope, IndexProfile
from rag_core.workers.runtime import WorkerConfig, run_job

pytestmark = [pytest.mark.integration, pytest.mark.asyncio]
BUCKET = "rag-core-storage-test"
SECRET_DIR = Path("/run/secrets")


class WireProxy:
    """Faults on real HTTP I/O, always forward actual writes before observing/gating."""

    def __init__(self, target: str, *, gate: bool = False, fail_after: int | None = None,
                 corrupt: bool = False) -> None:
        self.writes = 0
        self.written = threading.Event()
        self.release = threading.Event()
        self.fail_after, self.corrupt = fail_after, corrupt
        self.point: str | None = None
        if not gate:
            self.release.set()
        proxy = self

        class Handler(BaseHTTPRequestHandler):
            def log_message(self, *_args: object) -> None:
                pass

            def forward(self) -> None:
                payload = self.rfile.read(int(self.headers.get("Content-Length", "0")))
                write = self.command == "PUT" and self.path.split("?")[0].endswith("/points")
                if write and proxy.fail_after is not None and proxy.writes >= proxy.fail_after:
                    self.send_response(503)
                    self.end_headers()
                    self.wfile.write(b'{"status":{"error":"injected network outage"}}')
                    return
                with httpx.Client(base_url=target, trust_env=False, timeout=20) as client:
                    response = client.request(self.command, self.path, content=payload,
                                              headers={"Content-Type": "application/json"})
                    if write:
                        assert response.status_code == 200
                        proxy.writes += 1
                        proxy.point = json.loads(payload)["points"][0]["id"]
                        proxy.written.set()
                        if proxy.corrupt:
                            # Real server deletion after acknowledged write, before reconciliation.
                            path = self.path.split("?")[0] + "/delete?wait=true"
                            deletion = client.post(path, json={"points": [proxy.point]})
                            assert deletion.status_code == 200
                        proxy.release.wait(30)
                    self.send_response(response.status_code)
                    self.send_header("Content-Type", "application/json")
                    self.send_header("Content-Length", str(len(response.content)))
                    self.end_headers()
                    with suppress(BrokenPipeError, ConnectionResetError):
                        self.wfile.write(response.content)

            do_GET = do_POST = do_PUT = do_DELETE = forward

        self.server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()
        self.url = f"http://127.0.0.1:{self.server.server_port}"

    def close(self) -> None:
        self.release.set()
        self.server.shutdown()
        self.server.server_close()
        self.thread.join(5)


async def eventually(check: Callable[[], Any], *, timeout: float = 90) -> Any:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        value = check()
        if hasattr(value, "__await__"):
            value = await value
        if value:
            return value
        await asyncio.sleep(0.1)
    pytest.fail("Real service observation deadline exceeded")


@dataclass
class Stack:
    engine: AsyncEngine
    repo: PostgresRegistrationRepository
    worker_repo: PostgresIngestionRepository
    principal: Principal
    session_id: UUID
    config: WorkerConfig
    config_path: Path
    uploader: Any
    dispatcher: OutboxDispatcher
    broker: redis.Redis
    processes: list[subprocess.Popen[bytes]]
    tmp: Path
    pg_url: URL

    async def register(self, path: Path, mime: str, *, versioned: bool = True) -> RegisteredUpload:
        key = "allowed/t19-" + uuid4().hex + path.suffix
        data = path.read_bytes()
        version = self.uploader.put_object(Bucket=BUCKET, Key=key, Body=data)["VersionId"]
        source = {"storage_alias": "fixture", "bucket": BUCKET, "key": key,
                  "sha256": hashlib.sha256(data).hexdigest()}
        if versioned:
            source["version_id"] = version
        body = DocumentRegisterRequest.model_validate({"source": source, "filename": path.name,
            "content_type": mime, "external_upload_id": uuid4().hex})
        return await self.repo.register(self.principal, self.session_id, uuid4().hex, body)

    def start_worker(self, config: WorkerConfig | None = None) -> subprocess.Popen[bytes]:
        selected = self.tmp / ("worker-" + uuid4().hex + ".json")
        selected.write_text((config or self.config).model_dump_json())
        env = {**os.environ, "WORKER_CONFIG_FILE": str(selected),
               "DATABASE_URL": self.pg_url.render_as_string(hide_password=False)}
        # Environment contains the private DSN only in child process, never print it.
        env.pop("DATABASE_PASSWORD_FILE", None)
        process = subprocess.Popen([
            "celery", "-A", "rag_core.workers.celery:app", "worker", "--pool=prefork",
            "--concurrency=1", "--prefetch-multiplier=1", "-Q", "rag_core_ingestion",
            "--loglevel=ERROR", "--without-gossip", "--without-mingle", "--without-heartbeat",
        ], env=env, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, start_new_session=True)
        self.processes.append(process)
        return process

    async def job_ready(self, job_id: UUID) -> bool:
        return (await self.repo.get_job(self.principal, job_id)).state == "ready"

    async def run(self, upload: RegisteredUpload, config: WorkerConfig | None = None) -> str:
        # Factory uses its own engine and the real inference HTTP client.
        old_url = os.environ.get("DATABASE_URL")
        os.environ["DATABASE_URL"] = self.pg_url.render_as_string(hide_password=False)
        old_file = os.environ.pop("DATABASE_PASSWORD_FILE", None)
        try:
            return await run_job(upload.job.job_id, config or self.config)
        finally:
            if old_url is None:
                os.environ.pop("DATABASE_URL", None)
            else:
                os.environ["DATABASE_URL"] = old_url
            if old_file is not None:
                os.environ["DATABASE_PASSWORD_FILE"] = old_file

    async def report(self, upload: RegisteredUpload) -> None:
        async with self.engine.connect() as conn:
            row = (await conn.execute(text("""
                SELECT g.state,g.chunk_count,g.manifest_sha256,v.active_generation_id
                FROM document_versions v JOIN index_generations g
                    ON g.generation_id=v.active_generation_id WHERE v.version_id=:version
            """), {"version": upload.version_id})).mappings().one()
            chunks = (await conn.execute(text("""
                SELECT source_map,text FROM chunks WHERE generation_id=:gen ORDER BY ordinal
            """), {"gen": row["active_generation_id"]})).mappings().all()
        assert row["state"] == "ready" and row["chunk_count"] == len(chunks) > 0
        assert all(c["source_map"]["source"]["sha256"] for c in chunks)
        print(json.dumps({"job": str(upload.job.job_id), "state": "ready", "chunks": len(chunks),
            "locator": chunks[0]["source_map"]["segments"][0]["locator"],
            "excerpt": chunks[0]["text"][:180]}, ensure_ascii=False))


@pytest_asyncio.fixture
async def stack(pg_url: URL, tmp_path: Path) -> AsyncIterator[Stack]:
    assert os.environ.get("RAG_TEST_COMPOSE_ISOLATED") == "1", "Run isolated T19 Compose acceptance"
    model_url, qdrant_url = os.environ["INFERENCE_URL"], os.environ["QDRANT_URL"]
    async with httpx.AsyncClient(trust_env=False) as client:
        response = await client.get(model_url + "/health/ready", timeout=5)
        response.raise_for_status()
        info = response.json()
    assert info["model_instances"] == 2 and info["inference_processes"] == 1 and info["device"] == "cpu"
    principal = Principal("t19-" + uuid4().hex, "owner")
    registry = StorageRegistry(locations=(StorageLocation(
        app_id=principal.app_id, alias="fixture", endpoint="http://minio:9000", allow_compose_http=True,
        bucket=BUCKET, region="us-east-1", prefix="allowed/",
        access_key_file=SECRET_DIR / "minio_reader_user",
        secret_key_file=SECRET_DIR / "minio_reader_password",
    ),))
    storage_config = tmp_path / "storage.json"
    storage_config.write_text(registry.model_dump_json())
    config = WorkerConfig(model_fingerprint=info["fingerprint"], inference_url=model_url,
        qdrant_url=qdrant_url, storage_config=storage_config, temp_root=tmp_path / "worker-temp",
        lease_seconds=10, heartbeat_seconds=1)
    config_path = tmp_path / "worker.json"
    config_path.write_text(config.model_dump_json())
    temp = tmp_path / "registration-temp"
    temp.mkdir()
    engine = build_engine(pg_url)
    repo = PostgresRegistrationRepository(engine, S3StorageReader(registry, temp_dir=temp))
    uploader = boto3.client("s3", endpoint_url="http://minio:9000", region_name="us-east-1",
        aws_access_key_id=(SECRET_DIR / "minio_uploader_user").read_text().strip(),
        aws_secret_access_key=(SECRET_DIR / "minio_uploader_password").read_text().strip())
    session = await repo.sessions.create_session(principal, uuid4().hex)
    broker = redis.Redis.from_url(os.environ["REDIS_URL"], socket_timeout=3)
    assert broker.ping()
    f = Stack(engine, repo, PostgresIngestionRepository(engine, lease_seconds=10), principal,
        session.session_id, config, config_path, uploader,
        OutboxDispatcher(engine, CeleryJobPublisher(os.environ["REDIS_URL"])), broker, [], tmp_path, pg_url)
    started = time.monotonic()
    try:
        yield f
    finally:
        for process in f.processes:
            if process.poll() is None:
                os.killpg(process.pid, signal.SIGTERM)
                try:
                    await asyncio.to_thread(process.wait, timeout=5)
                except subprocess.TimeoutExpired:
                    os.killpg(process.pid, signal.SIGKILL)
                    process.wait()
        await engine.dispose()
        broker.close()
        # Own random module DB is dropped by pg_url; MinIO test project is disposable.
        assert not temp.exists() or not list(temp.iterdir())
        peak = Path("/sys/fs/cgroup/memory.peak")
        print(f"REAL ingestion case elapsed={time.monotonic()-started:.3f}s "
              f"test_container_peak_mib={int(peak.read_text())/1024**2:.3f} shared_model_pid={info['pid']}")


def source_text(stack: Stack) -> Path:
    path = stack.tmp / "source.txt"
    phrase = "Quarterly revenue reached 120 million VND. Doanh thu đạt 120 triệu đồng.\n"
    path.write_text(phrase * 2, encoding="utf-8")
    return path


async def test_scan_xlsx_text_real_celery_and_source_locators(stack: Stack) -> None:
    f = stack
    font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 48)
    pdf = f.tmp / "scan.pdf"
    with Image.new("RGB", (1800, 420), "white") as bitmap:
        ImageDraw.Draw(bitmap).text((65, 100), "Doanh thu quý một đạt 120 triệu đồng.", font=font,
                                   fill="black")
        bitmap.save(pdf, "PDF", resolution=216)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Revenue VND"
    sheet.append(["Quarter", "Revenue (million VND)"])
    sheet.append(["Q1", 120])
    sheet.append(["Q2", 150])
    xlsx = f.tmp / "revenue.xlsx"
    workbook.save(xlsx)
    txt = source_text(f)
    inputs = [(pdf, "application/pdf"), (xlsx,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"), (txt, "text/plain")]
    before = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p, _ in inputs}
    uploads = [await f.register(p, mime) for p, mime in inputs]
    f.start_worker()
    for upload in uploads:
        assert await f.dispatcher.dispatch_one()
        await eventually(lambda upload=upload: f.job_ready(upload.job.job_id))
        assert await f.run(upload) == "not_claimed"
        await f.report(upload)
    snapshot = await f.repo.sessions.resolve_scope(f.principal, f.session_id)
    assert len(snapshot.pairs) == 3
    async with f.engine.connect() as conn:
        maps = (await conn.execute(text("SELECT source_map,text FROM chunks WHERE app_id=:app"),
                    {"app": f.principal.app_id})).mappings().all()
    pdf_maps = [m for m in maps if m["source_map"]["segments"][0]["locator"]["kind"] == "pdf"]
    assert pdf_maps and "Doanh thu quý một đạt 120 triệu đồng." in pdf_maps[0]["text"]
    assert all(s["locator"]["page"] == 1 for m in pdf_maps for s in m["source_map"]["segments"])
    xlsx_maps = [m for m in maps if m["source_map"]["segments"][0]["locator"]["kind"] == "xlsx"]
    assert xlsx_maps and any(s["locator"]["cell_range"] == "B2" for m in xlsx_maps
                            for s in m["source_map"]["segments"])
    assert all(hashlib.sha256(p.read_bytes()).hexdigest() == sha for p, sha in before.items())
    # App object bytes stay identical, reader has never PUT/DELETE permissions.
    async with f.engine.connect() as conn:
        objects = (await conn.execute(text("SELECT object_key,source_sha256,source_version_id "
            "FROM documents d JOIN document_versions v USING(app_id,owner_id,document_id) "
            "WHERE app_id=:app"), {"app": f.principal.app_id})).mappings().all()
    for obj in objects:
        data = f.uploader.get_object(Bucket=BUCKET, Key=obj["object_key"],
                                    VersionId=obj["source_version_id"])["Body"].read()
        assert hashlib.sha256(data).hexdigest() == obj["source_sha256"]


async def test_duplicate_pending_generation_invisible(stack: Stack) -> None:
    f = stack
    upload = await f.register(source_text(f), "text/plain")
    proxy = WireProxy(f.config.qdrant_url, gate=True)
    config = f.config.model_copy(update={"qdrant_url": proxy.url})
    task = asyncio.create_task(f.run(upload, config))
    try:
        await eventually(proxy.written.is_set)
        assert await f.run(upload) == "not_claimed"
        assert (await f.repo.get_job(f.principal, upload.job.job_id)).attempts == 1
        async with f.engine.connect() as conn:
            before = (await conn.execute(text("SELECT lease_until FROM ingestion_jobs WHERE job_id=:job"),
                                         {"job": upload.job.job_id})).scalar_one()
        await asyncio.sleep(2.1)
        async with f.engine.connect() as conn:
            after = (await conn.execute(text("SELECT lease_until FROM ingestion_jobs WHERE job_id=:job"),
                                        {"job": upload.job.job_id})).scalar_one()
        assert after > before  # Heartbeats renew during the actual blocked Qdrant ACK.
        with pytest.raises(ScopeError, match="documents_not_ready"):
            await f.repo.sessions.resolve_scope(f.principal, f.session_id)
        proxy.release.set()
        assert await task == "ready"
        await f.report(upload)
    finally:
        proxy.close()
        await asyncio.gather(task, return_exceptions=True)


async def test_partial_qdrant_failure_retry_keeps_old_generation(stack: Stack) -> None:
    f = stack
    path = f.tmp / "multi-batch.xlsx"
    workbook = Workbook()
    workbook.remove(workbook.active)
    for number in range(34):
        sheet = workbook.create_sheet(f"Quarter {number}")
        sheet.append(["Quarter", "Revenue (million VND)"])
        sheet.append([str(number), 120 + number])
    workbook.save(path)
    upload = await f.register(path,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    print("REAL reindex: initial multi-batch source ingestion")
    assert await f.run(upload) == "ready"
    old = await f.repo.sessions.resolve_scope(f.principal, f.session_id)
    await f.worker_repo.request_reindex(f.principal, upload.job.job_id)
    proxy = WireProxy(f.config.qdrant_url, fail_after=1)
    try:
        print("REAL reindex: acknowledged first batch, next write forced HTTP503")
        assert await f.run(upload, f.config.model_copy(update={"qdrant_url": proxy.url})) == "vector_dependency_unavailable"
        assert proxy.writes == 1
        assert (await f.repo.sessions.resolve_scope(f.principal, f.session_id)).pairs == old.pairs
        async with f.engine.connect() as conn:
            failed = (await conn.execute(text("SELECT generation_id FROM index_generations "
                "WHERE job_id=:job AND state='failed'"), {"job": upload.job.job_id})).scalar_one()
        client = AsyncQdrantClient(url=f.config.qdrant_url, trust_env=False)
        profile = IndexProfile(f.config.model_fingerprint)
        from rag_core.adapters.persistence.vector_generations import PostgresGenerationAuthority
        from rag_core.adapters.vectors.qdrant import QdrantVectorRepository
        vectors = QdrantVectorRepository(client, profile, f.repo.sessions,
                                         PostgresGenerationAuthority(f.engine))
        failed_scope = GenerationScope(f.principal, VersionGeneration(
            upload.document_id, upload.version_id, failed))
        assert await vectors.count_generation(failed_scope) == 32
        await vectors.cleanup_generation(failed_scope)
        assert await vectors.count_generation(failed_scope) == 0
        assert await vectors.count_generation(GenerationScope(f.principal, old.pairs[0])) > 32
        await client.close()
        print("REAL reindex: retry new generation with retained old publication")
        assert await f.run(upload) == "ready"
        new = await f.repo.sessions.resolve_scope(f.principal, f.session_id)
        assert new.pairs != old.pairs
        with pytest.raises(ScopeError, match="session_scope_changed"):
            await f.repo.sessions.validate_snapshot(old)
        assert (await f.repo.get_job(f.principal, upload.job.job_id)).attempts == 2
        await f.report(upload)
    finally:
        proxy.close()


async def test_changed_source_fails_without_publication(stack: Stack) -> None:
    f = stack
    upload = await f.register(source_text(f), "text/plain", versioned=False)
    async with f.engine.connect() as conn:
        key = (await conn.execute(text("SELECT object_key FROM documents WHERE document_id=:doc"),
                           {"doc": upload.document_id})).scalar_one()
    f.uploader.put_object(Bucket=BUCKET, Key=key, Body=b"App changed source after registration")
    assert await f.run(upload) == "source_changed"
    job = await f.repo.get_job(f.principal, upload.job.job_id)
    assert job.state == "failed" and job.error_code == "source_changed"
    with pytest.raises(ScopeError, match="documents_not_ready"):
        await f.repo.sessions.resolve_scope(f.principal, f.session_id)


@pytest.mark.parametrize("action", ["detach", "delete"])
async def test_visibility_detach_delete_no_resurrection(stack: Stack, action: str) -> None:
    f = stack
    upload = await f.register(source_text(f), "text/plain")
    proxy = WireProxy(f.config.qdrant_url, gate=True)
    task = asyncio.create_task(f.run(upload, f.config.model_copy(update={"qdrant_url": proxy.url})))
    try:
        await eventually(proxy.written.is_set)
        if action == "detach":
            await f.repo.detach_document(f.principal, f.session_id, upload.document_id)
        else:
            await f.repo.sessions.delete_session(f.principal, f.session_id)
        proxy.release.set()
        assert await task == "ingestion_cancelled"
        assert (await f.repo.get_job(f.principal, upload.job.job_id)).state == "cancelled"
        assert await f.run(upload) == "not_claimed"
        with pytest.raises(ScopeError):
            await f.repo.sessions.resolve_scope(f.principal, f.session_id)
        async with f.engine.connect() as conn:
            assert (await conn.execute(text("SELECT count(*) FROM session_documents "
                "WHERE app_id=:app AND status='attached'"), {"app": f.principal.app_id})).scalar_one() == 0
            assert (await conn.execute(text("SELECT count(*) FROM chunks WHERE app_id=:app"),
                {"app": f.principal.app_id})).scalar_one() > 0  # Retained derivatives.
            generation = (await conn.execute(text("SELECT generation_id FROM index_generations "
                "WHERE job_id=:job"), {"job": upload.job.job_id})).scalar_one()
            obj = (await conn.execute(text("SELECT d.object_key,v.source_sha256,v.source_version_id "
                "FROM documents d JOIN document_versions v USING(app_id,owner_id,document_id) "
                "WHERE v.version_id=:version"), {"version": upload.version_id})).mappings().one()
        data = f.uploader.get_object(Bucket=BUCKET, Key=obj["object_key"],
                                    VersionId=obj["source_version_id"])["Body"].read()
        assert hashlib.sha256(data).hexdigest() == obj["source_sha256"]
        from rag_core.adapters.persistence.vector_generations import PostgresGenerationAuthority
        from rag_core.adapters.vectors.qdrant import QdrantVectorRepository
        client = AsyncQdrantClient(url=f.config.qdrant_url, trust_env=False)
        try:
            vectors = QdrantVectorRepository(client, IndexProfile(f.config.model_fingerprint),
                f.repo.sessions, PostgresGenerationAuthority(f.engine))
            scope = GenerationScope(f.principal, VersionGeneration(upload.document_id,
                upload.version_id, generation))
            assert await vectors.count_generation(scope) == 1
        finally:
            await client.close()
        print(f"REAL {action} during indexing: cancelled, attached_links=0, chunks retained")
    finally:
        proxy.close()
        await asyncio.gather(task, return_exceptions=True)


async def test_worker_kill_resume_and_expired_fence(stack: Stack) -> None:
    f = stack
    upload = await f.register(source_text(f), "text/plain")
    proxy = WireProxy(f.config.qdrant_url, gate=True)
    try:
        process = f.start_worker(f.config.model_copy(update={"qdrant_url": proxy.url}))
        while await f.dispatcher.dispatch_one():
            pass
        await eventually(proxy.written.is_set)
        os.killpg(process.pid, signal.SIGKILL)
        await asyncio.to_thread(process.wait, timeout=5)
        assert process.returncode == -signal.SIGKILL
        assert await f.run(upload) == "not_claimed"  # Live lease, duplicate cannot steal it.
        proxy.release.set()
        await asyncio.sleep(11)  # Actual configured lease expiration, no clock/mock substitution.
        assert await f.worker_repo.recover_one()
        job = await f.repo.get_job(f.principal, upload.job.job_id)
        assert job.state == "queued" and job.attempts == 1
        async with f.engine.connect() as conn:
            old = (await conn.execute(text("SELECT * FROM index_generations WHERE job_id=:job"),
                {"job": upload.job.job_id})).mappings().one()
        assert old["state"] == "failed"
        from rag_core.domain.ingestion import IngestionLease
        from rag_core.ports.storage import SourceObject
        stale = IngestionLease(upload.job.job_id, f.principal, f.session_id,
            GenerationScope(f.principal, VersionGeneration(upload.document_id,upload.version_id,
                                                           old["generation_id"])),
            old["lease_owner"], SourceObject("fixture", BUCKET, "allowed/stale", sha256="a"*64),
            "source.txt", "text/plain", 1)
        with pytest.raises(IngestionError, match="ingestion_lease_lost"):
            await f.worker_repo.publish(stale, 1, "a"*64, "stale")
        f.start_worker()
        await asyncio.sleep(2.1)  # Durable bounded backoff for attempt1.
        while await f.dispatcher.dispatch_one():
            pass
        await eventually(lambda: f.job_ready(upload.job.job_id))
        assert (await f.repo.get_job(f.principal, upload.job.job_id)).attempts == 2
        await f.report(upload)
        print(f"REAL Celery worker killed exit={process.returncode}; expired lease fenced, resume attempts=2")
    finally:
        proxy.close()


async def test_incomplete_vector_manifest_never_ready(stack: Stack) -> None:
    f = stack
    upload = await f.register(source_text(f), "text/plain")
    proxy = WireProxy(f.config.qdrant_url, corrupt=True)
    try:
        assert await f.run(upload, f.config.model_copy(update={"qdrant_url": proxy.url})) == "vector_manifest_mismatch"
        assert (await f.repo.get_job(f.principal, upload.job.job_id)).state == "failed"
        with pytest.raises(ScopeError, match="documents_not_ready"):
            await f.repo.sessions.resolve_scope(f.principal, f.session_id)
    finally:
        proxy.close()


async def test_partial_pdf_refused_before_embedding(stack: Stack) -> None:
    from pypdf import PdfReader, PdfWriter

    f = stack
    path = source_text(f)
    # Use actual image-PDF plus genuinely blank second page, never mocked OCR output.
    pdf = f.tmp / "partial.pdf"
    font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 48)
    with Image.new("RGB", (1800, 420), "white") as bitmap:
        ImageDraw.Draw(bitmap).text((65, 100), "Revenue was 120 million VND.", font=font, fill="black")
        bitmap.save(pdf, "PDF", resolution=216)
    writer = PdfWriter()
    writer.add_page(PdfReader(pdf).pages[0])
    writer.add_blank_page(width=600, height=150)
    with pdf.open("wb") as output:
        writer.write(output)
    upload = await f.register(pdf, "application/pdf")
    assert path.exists()
    assert await f.run(upload) == "partial_extraction"
    assert (await f.repo.get_job(f.principal, upload.job.job_id)).state == "failed"


async def test_retry_budget_and_durable_outbox(stack: Stack) -> None:
    from rag_core.adapters.persistence.registrations import RegistrationError

    f = stack
    upload = await f.register(source_text(f), "text/plain")
    proxy = WireProxy(f.config.qdrant_url, fail_after=0)
    try:
        for attempt in range(1, 4):
            assert await f.run(upload, f.config.model_copy(update={"qdrant_url": proxy.url})) == "vector_dependency_unavailable"
            job = await f.repo.get_job(f.principal, upload.job.job_id)
            assert job.attempts == attempt
            assert job.state == ("failed" if attempt == 3 else "queued")
        with pytest.raises(RegistrationError, match="job_not_retryable"):
            await f.repo.retry_job(f.principal, upload.job.job_id)
        async with f.engine.connect() as conn:
            count = (await conn.execute(text("SELECT count(*) FROM outbox_events WHERE job_id=:job"),
                                       {"job": upload.job.job_id})).scalar_one()
        assert count == 3  # Original registration + two retries, no endless retry loop.
        print("REAL retry budget: attempts=3, failed, durable events=3, further retry denied")
    finally:
        proxy.close()


@pytest.mark.parametrize("column", ["text", "source_map"])
async def test_pg_manifest_tamper_never_published(stack: Stack, column: str) -> None:
    f = stack
    upload = await f.register(source_text(f), "text/plain")
    proxy = WireProxy(f.config.qdrant_url, gate=True)
    task = asyncio.create_task(f.run(upload, f.config.model_copy(update={"qdrant_url": proxy.url})))
    try:
        await eventually(proxy.written.is_set)
        # Actual database corruption after staging, before publication; immutable manifest catches it.
        async with f.engine.begin() as conn:
            change = "text=text || ' altered'" if column == "text" else (
                "source_map=jsonb_set(source_map,'{heading_path}','[\"altered\"]'::jsonb)")
            await conn.execute(text(f"UPDATE chunks SET {change} WHERE version_id=:version"),
                               {"version": upload.version_id})
        proxy.release.set()
        assert await task == "ingestion_manifest_mismatch"
        assert (await f.repo.get_job(f.principal, upload.job.job_id)).state == "failed"
        with pytest.raises(ScopeError, match="documents_not_ready"):
            await f.repo.sessions.resolve_scope(f.principal, f.session_id)
    finally:
        proxy.close()
        await asyncio.gather(task, return_exceptions=True)
