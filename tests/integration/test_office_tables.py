"""Actual Office/CSV parsers on synthetic fixtures, with source and execution canaries."""

import csv
import hashlib
import io
import threading
import zipfile
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from uuid import UUID
from xml.etree import ElementTree as ET

import pytest
from docx import Document
from lxml import etree
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches

from rag_core.adapters.parsers import ParserRegistry
from rag_core.adapters.parsers.registry import FORMATS
from rag_core.contracts.v1 import XlsxLocator
from rag_core.domain.documents import ParsedDocument, ParseError, ParserLimits, SourceIdentity

pytestmark = pytest.mark.integration
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def parse(
    path: Path, *, limits: ParserLimits | None = None, content_type: str | None = None
) -> ParsedDocument:
    original = path.read_bytes()
    sandbox = path.parent / "sandbox"
    try:
        return ParserRegistry(sandbox, limits).parse(
            path,
            filename=path.name,
            content_type=content_type or FORMATS[path.suffix][1],
            source=SourceIdentity(
                document_id=UUID(int=1),
                version_id=UUID(int=2),
                sha256=hashlib.sha256(original).hexdigest(),
            ),
        )
    finally:
        assert path.read_bytes() == original
        assert not sandbox.exists() or list(sandbox.iterdir()) == []


def workbook(path: Path) -> None:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Doanh thu"
    sheet.append(["Báo cáo EN/VI", None, None])
    sheet.merge_cells("A1:C1")
    sheet.append(["Quý", "Doanh thu (triệu đồng)", "Tỷ lệ"])
    sheet.append(["Q1", 120, 0.25])
    sheet["C3"].number_format = "0%"
    sheet.append(["Q2", "=B3*2", "=C3*2"])
    other = book.create_sheet("Expenses USD")
    other.append(["Item", "Cost (USD)"])
    other.append(["Hosting", 45])
    book.save(path)
    # A real stored Excel cache is source fixture data, not a mocked parser result.
    with zipfile.ZipFile(path) as package:
        members = {name: package.read(name) for name in package.namelist()}
    root = ET.fromstring(members["xl/worksheets/sheet1.xml"])
    cell = root.find("s:sheetData/s:row/s:c[@r='B4']", NS)
    assert cell is not None
    cached = cell.find("s:v", NS)
    assert cached is not None
    cached.text = "240"
    members["xl/worksheets/sheet1.xml"] = ET.tostring(root)
    write_package(path, members)


def presentation(path: Path) -> None:
    deck = Presentation()
    for number in (1, 2):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        box.text_frame.text = f"Báo cáo slide {number} / Revenue"
        box.text_frame.add_paragraph()  # blank counts in source block index
        box.text_frame.add_paragraph().text = "Doanh thu 120 triệu đồng."
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(5), Inches(2)).table
        for row, values in enumerate((("Quý", "Amount (triệu đồng)"), ("Q1", "120"))):
            for column, value in enumerate(values):
                table.cell(row, column).text = value
        group = slide.shapes.add_group_shape()
        group.shapes.add_textbox(Inches(1), Inches(5), Inches(3), Inches(1)).text = "Grouped EN/VI"
        group.shapes[0].text_frame.add_paragraph()
        group.shapes.add_textbox(
            Inches(1), Inches(6), Inches(3), Inches(1)
        ).text = "Second group child"
    deck.save(path)


def write_package(path: Path, members: dict[str, bytes]) -> None:
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in members.items():
            archive.writestr(name, data)


def members(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def make_office(tmp_path: Path, format_: str) -> Path:
    path = tmp_path / f"report.{format_}"
    (workbook if format_ == "xlsx" else presentation)(path)
    return path


def test_xlsx_multisheet_merged_headers_units_and_cell_round_trip(tmp_path: Path) -> None:
    path = make_office(tmp_path, "xlsx")
    result = parse(path)
    assert result.format == "xlsx" and result.page_count is None
    book = load_workbook(path, data_only=False)
    cache = load_workbook(path, data_only=True)
    assert {block.locator.sheet for block in result.blocks} == {"Doanh thu", "Expenses USD"}
    for block in result.blocks:
        locator = block.locator
        assert locator.kind == "xlsx"
        assert XlsxLocator.model_validate_json(locator.model_dump_json()) == locator
        sheet = book[locator.sheet]
        expected = tuple(cell.coordinate for row in sheet[locator.cell_range] for cell in row)
        assert tuple(cell.reference for cell in block.cells) == expected
        assert tuple(cell.value for cell in block.cells) == block.rows[0]
        assert block.source_part and block.source_path
        raw = etree.fromstring(members(path)[block.source_part])
        source_row = raw.getroottree().xpath(block.source_path, namespaces=NS)[0]
        assert source_row.attrib["r"] == str(sheet[block.cells[0].reference].row)
        for cell in block.cells:
            original = sheet[cell.reference]
            if cell.formula:
                assert cell.formula == original.value
                if cell.value_origin == "formula_cache":
                    assert cell.value == str(cache[locator.sheet][cell.reference].value)
                else:
                    assert cache[locator.sheet][cell.reference].value is None
                    assert cell.value == "[formula cache unavailable]"
            else:
                assert cell.value == (str(original.value) if original.value is not None else "")
    q1 = next(
        block
        for block in result.blocks
        if block.locator.sheet == "Doanh thu" and block.cells[0].reference == "A3"
    )
    assert q1.table_headers[0] == ("Báo cáo EN/VI",) * 3
    assert q1.table_headers[1][1] == "Doanh thu (triệu đồng)"
    assert "Doanh thu (triệu đồng)" in q1.text and "0%" in q1.text and "0.25" in q1.text
    merged = result.blocks[0]
    assert merged.rows[0] == ("Báo cáo EN/VI", "", "")
    assert merged.cells[1].merged_origin == "A1"
    book.close()
    cache.close()
    print(
        "PASS actual XLSX: 2 sheets, merged A1:C1, headers/units/percent format, every cell/range round-trip; no PDF page"
    )


def test_xlsx_formula_cache_policy_never_recalculates(tmp_path: Path) -> None:
    result = parse(make_office(tmp_path, "xlsx"))
    formula = {
        cell.reference: cell for block in result.blocks for cell in block.cells if cell.formula
    }
    assert formula["B4"].formula == "=B3*2"
    assert formula["B4"].value == "240" and formula["B4"].value_origin == "formula_cache"
    assert formula["C4"].formula == "=C3*2"
    assert formula["C4"].value == "[formula cache unavailable]"
    assert formula["C4"].value_origin == "missing_formula_cache"
    assert set(result.warnings) == {"formula_cache_missing", "formula_cached_values_unverified"}
    assert "=B3*2" not in "\n".join(block.text for block in result.blocks)
    print(
        "PASS formula policy: stored cache240 retained, missing cache not inferred as0.5; formulas separate and inert"
    )


def test_xlsx_numeric_headers_merged_extent_blank_regions_and_false_dimension(
    tmp_path: Path,
) -> None:
    path = tmp_path / "context.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["Revenue (USD)", 2023, 2024])
    sheet.append(["EN/VI", 10, 20])
    sheet.append([])
    sheet.append(["Cost (VND)", None, None])
    sheet.merge_cells("A4:D4")
    sheet.append(["Chi phí", 30, 40, 50])
    book.save(path)
    package = members(path)
    xml = ET.fromstring(package["xl/worksheets/sheet1.xml"])
    dimension = xml.find("s:dimension", NS)
    assert dimension is not None
    dimension.attrib["ref"] = "A1:XFD1048576"
    package["xl/worksheets/sheet1.xml"] = ET.tostring(xml)
    write_package(path, package)
    result = parse(path)
    assert len(result.blocks) == 4
    assert result.blocks[1].table_headers == (("Revenue (USD)", "2023", "2024", ""),)
    assert result.blocks[2].locator.cell_range == "A4:D4"
    assert result.blocks[-1].table_headers == (("Cost (VND)",) * 4,)
    assert "Revenue (USD)" not in result.blocks[-1].text
    print(
        "PASS numeric year headers/units, full merged extent, blank-region reset; forged dimension ignored"
    )


def test_common_table_normalization_preserves_empty_cells_multiline_headers_units(
    tmp_path: Path,
) -> None:
    rows = (("Item", "Amount (USD)", "Note"), ("Đà Nẵng", "120", ""))
    doc = Document()
    table = doc.add_table(rows=2, cols=3)
    for n, row in enumerate(rows):
        for c, value in enumerate(row):
            table.cell(n, c).text = value
    docx = tmp_path / "table.docx"
    doc.save(docx)
    html = tmp_path / "table.html"
    html.write_text(
        "<html><table><tr><th>Item</th><th>Amount (USD)</th><th>Note</th></tr><tr><td>Đà Nẵng</td><td>120</td><td></td></tr></table></html>",
        encoding="utf-8",
    )
    md = tmp_path / "table.md"
    md.write_text(
        "| Item | Amount (USD) | Note |\n| --- | --- | --- |\n| Đà Nẵng | 120 | |", encoding="utf-8"
    )
    for path in (docx, html, md):
        block = next(block for block in parse(path).blocks if block.kind == "table")
        assert block.rows == rows and block.table_headers == rows[:1]
        assert "Amount (USD)" in block.text
    csv_path = tmp_path / "inert.csv"
    csv_path.write_text('Header,Note\n"=1+2","<p>EN/VI</p>"', encoding="utf-8")
    assert parse(csv_path).blocks[1].rows == (("=1+2", "<p>EN/VI</p>"),)
    print(
        "PASS shared DOCX/HTML/MD table rows/header/units/empty cells; CSV HTML/formula-looking fields remain inert data"
    )


@pytest.mark.parametrize("delimiter", [",", ";", "\t", "|"])
def test_csv_unicode_quoting_multiline_delimiters_record_round_trip(
    tmp_path: Path, delimiter: str
) -> None:
    rows = [
        ["Khoản mục", "Value (triệu đồng)", "Ghi chú"],
        ["Doanh thu", "120", 'EN, VI; | tab\t "quoted"\nsecond line'],
        ["=1+2", "", "Đà Nẵng"],
    ]
    stream = io.StringIO(newline="")
    csv.writer(stream, delimiter=delimiter, lineterminator="\r\n").writerows(rows)
    path = tmp_path / "data.csv"
    path.write_bytes(("\ufeff" + stream.getvalue()).encode("utf-8"))
    result = parse(path)
    original = list(
        csv.reader(
            io.StringIO(path.read_text(encoding="utf-8-sig"), newline=""), delimiter=delimiter
        )
    )
    assert len(result.blocks) == 3
    for index, block in enumerate(result.blocks, 1):
        locator = block.locator
        assert locator.kind == "csv" and locator.row_start == locator.row_end == index
        assert locator.columns == rows[0]
        assert block.rows == (tuple(original[index - 1]),)
        assert block.table_headers == (tuple(rows[0]),)
        assert "triệu đồng" in block.text
        assert type(locator).model_validate_json(locator.model_dump_json()) == locator
    assert result.blocks[-1].rows[0][0] == "=1+2"
    print(
        f"PASS actual CSV delimiter{delimiter!r}: Unicode/BOM/CRLF/escaped quotes/multiline; logical records1-3, formula text unchanged"
    )


def test_pptx_multislide_table_group_xml_locator_round_trip(tmp_path: Path) -> None:
    path = make_office(tmp_path, "pptx")
    result = parse(path)
    deck = Presentation(path)
    package = members(path)
    assert result.page_count is None
    assert {block.locator.slide for block in result.blocks} == {1, 2}
    for block in result.blocks:
        locator = block.locator
        assert locator.kind == "pptx" and locator.shape is not None and locator.block is not None
        slide = deck.slides[locator.slide - 1]
        assert block.source_part == str(slide.part.partname).lstrip("/")
        assert block.source_path
        raw = etree.fromstring(package[block.source_part])
        element = raw.getroottree().xpath(block.source_path, namespaces=raw.nsmap)[0]
        shape = slide.shapes[locator.shape]
        if block.kind == "table":
            rows = tuple(tuple(cell.text for cell in row.cells) for row in shape.table.rows)
            assert block.rows == rows and block.table_headers == rows[:1]
            assert "triệu đồng" in block.text
        else:
            assert "".join(element.itertext()) == block.text
        assert type(locator).model_validate_json(locator.model_dump_json()) == locator
    assert any(block.locator.block == 2 for block in result.blocks)
    assert sum(block.text == "Grouped EN/VI" for block in result.blocks) == 2
    for slide in (1, 2):
        grouped = [
            block.locator.block
            for block in result.blocks
            if block.locator.slide == slide and block.locator.shape == 2
        ]
        assert grouped == [0, 2] and len(grouped) == len(set(grouped))
    print(
        "PASS actual PPTX: slides1/2, blank paragraph indices, tables/units, group paths and original XML round-trip"
    )


@pytest.mark.parametrize(
    "payload", [b"a,b\n1,2,3\n", b'a,b\n1,"unterminated', b"\xff\xfe\x00", b"a,b\n\n"]
)
def test_csv_invalid_encoding_shape_or_quotes_are_explicit_errors(
    tmp_path: Path, payload: bytes
) -> None:
    path = tmp_path / "bad.csv"
    path.write_bytes(payload)
    with pytest.raises(ParseError, match=r"invalid_csv|invalid_encoding"):
        parse(path)


@pytest.mark.parametrize("format_", ["xlsx", "pptx"])
def test_safety_office_archive_bombs_macros_traversal_entities(
    tmp_path: Path, format_: str
) -> None:
    path = make_office(tmp_path, format_)
    source = members(path)
    for limits in (
        ParserLimits(max_archive_entries=1),
        ParserLimits(max_expanded_bytes=100),
        ParserLimits(max_compression_ratio=1),
    ):
        with pytest.raises(ParseError, match="archive_limit"):
            parse(path, limits=limits)
    for name, value in (
        ("../outside.txt", b"unsafe"),
        ("payload/vbaProject.bin", b"macro canary"),
        ("evil.xml", b'<!DOCTYPE x [<!ENTITY e "bomb">]><x>&e;</x>'),
        ("bomb.bin", b"0" * 1_000_000),
    ):
        write_package(path, {**source, name: value})
        with pytest.raises(ParseError, match=r"unsafe_archive|corrupt_document|archive_limit"):
            parse(path)
    assert not (tmp_path / "outside.txt").exists()
    print(
        f"PASS {format_} actual ZIP entry/expanded/ratio bombs, traversal, VBA and XML entities refused; no extraction, source preserved"
    )


@pytest.mark.parametrize("format_", ["xlsx", "pptx"])
def test_safety_external_relationships_never_fetch_or_execute(tmp_path: Path, format_: str) -> None:
    calls: list[str] = []

    class Canary(BaseHTTPRequestHandler):
        def do_GET(self) -> None:
            calls.append(self.path)
            self.send_response(200)
            self.end_headers()

        def log_message(self, *args: object) -> None:
            pass

    server = ThreadingHTTPServer(("127.0.0.1", 0), Canary)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        path = make_office(tmp_path, format_)
        url = f"http://127.0.0.1:{server.server_port}/execute"
        if format_ == "xlsx":
            book = load_workbook(path)
            sheet = book.worksheets[0]
            sheet["B5"] = f'=WEBSERVICE("{url}")'
            sheet["C5"] = "External hyperlink"
            sheet["C5"].hyperlink = url
            book.save(path)
            book.close()
        else:
            deck = Presentation(path)
            deck.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].hyperlink.address = url
            deck.save(path)
        package = members(path)
        relationship = (
            "xl/_rels/workbook.xml.rels"
            if format_ == "xlsx"
            else "ppt/slides/_rels/slide1.xml.rels"
        )
        root = ET.fromstring(package[relationship])
        ET.SubElement(
            root,
            "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
            {
                "Id": "rIdCanary",
                "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                "Target": url,
                "TargetMode": "External",
            },
        )
        package[relationship] = ET.tostring(root)
        write_package(path, package)
        result = parse(path)
        assert "external_relationships_ignored" in result.warnings
        assert calls == []
    finally:
        server.shutdown()
        server.server_close()
        thread.join()
    print(
        f"PASS {format_} actual loopback HTTP execution canary:0 requests; external relations inert"
    )


def test_safety_csv_single_column_empty_and_cell_text_limits(tmp_path: Path) -> None:
    path = tmp_path / "single.csv"
    path.write_bytes('Name\nĐà Nẵng\n"Line one\nLine two"'.encode())
    result = parse(path)
    assert result.blocks[2].rows == (("Line one\nLine two",),)
    assert result.blocks[2].locator.row_start == 3
    with pytest.raises(ParseError, match="table_limit"):
        parse(path, limits=ParserLimits(max_table_cells=2))
    with pytest.raises(ParseError, match="extraction_limit"):
        parse(path, limits=ParserLimits(max_text_chars=5))
    path.write_bytes(b"")
    with pytest.raises(ParseError, match="empty_extraction"):
        parse(path)


def test_empty_csv_header_and_merged_formula_header_keep_source_data_distinct(
    tmp_path: Path,
) -> None:
    csv_path = tmp_path / "empty-header.csv"
    csv_path.write_bytes(b'""\nRevenue')
    csv_result = parse(csv_path)
    assert csv_result.blocks[0].rows == (("",),)
    assert csv_result.blocks[0].locator.columns == ["Column 1"]
    assert csv_result.blocks[1].rows == (("Revenue",),)
    csv_path.write_bytes(b'""')
    with pytest.raises(ParseError, match="empty_extraction"):
        parse(csv_path)
    path = tmp_path / "merged-formula.xlsx"
    book = Workbook()
    sheet = book.active
    sheet["A1"] = "=1+2"
    sheet.merge_cells("A1:B1")
    sheet.append(["Revenue", 120])
    book.save(path)
    result = parse(path)
    assert result.blocks[0].table_headers == (("[formula cache unavailable]",) * 2,)
    assert result.blocks[0].cells[0].formula == "=1+2"
    assert "=1+2" not in "\n".join(block.text for block in result.blocks)


def test_safety_xlsx_sparse_dimension_merged_and_cell_sheet_budgets(tmp_path: Path) -> None:
    path = make_office(tmp_path, "xlsx")
    for limits, error in (
        (ParserLimits(max_table_cells=2), "table_limit"),
        (ParserLimits(max_sheets=1), "sheet_limit"),
    ):
        with pytest.raises(ParseError, match=error):
            parse(path, limits=limits)
    package = members(path)
    root = ET.fromstring(package["xl/worksheets/sheet1.xml"])
    merge = root.find("s:mergeCells/s:mergeCell", NS)
    assert merge is not None
    merge.attrib["ref"] = "A1:XFD1048576"
    package["xl/worksheets/sheet1.xml"] = ET.tostring(root)
    write_package(path, package)
    with pytest.raises(ParseError, match="table_limit"):
        parse(path)
    print(
        "PASS XLSX pre-load sparse/merge area and sheet/cell budgets; huge merge never materialized"
    )


def test_safety_pptx_slide_table_and_result_budgets(tmp_path: Path) -> None:
    path = make_office(tmp_path, "pptx")
    for limits, error in (
        (ParserLimits(max_slides=1), "slide_limit"),
        (ParserLimits(max_table_cells=1), "table_limit"),
        (ParserLimits(max_result_bytes=100), "extraction_limit"),
    ):
        with pytest.raises(ParseError, match=error):
            parse(path, limits=limits)


@pytest.mark.parametrize("format_", ["xlsx", "pptx"])
def test_safety_office_mime_package_corrupt_encrypted_and_legacy(
    tmp_path: Path, format_: str
) -> None:
    path = make_office(tmp_path, format_)
    other = make_office(tmp_path, "pptx" if format_ == "xlsx" else "xlsx")
    path.write_bytes(other.read_bytes())
    with pytest.raises(ParseError, match="mime_mismatch"):
        parse(path)
    path.write_bytes(b"PK broken package")
    with pytest.raises(ParseError, match="corrupt_document"):
        parse(path)
    path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1 encrypted Office")
    with pytest.raises(ParseError, match="encrypted_document"):
        parse(path)
    legacy = tmp_path / ("old.xls" if format_ == "xlsx" else "old.ppt")
    legacy.write_bytes(path.read_bytes())
    with pytest.raises(ParseError, match="unsupported_format"):
        parse(legacy, content_type="application/octet-stream")
